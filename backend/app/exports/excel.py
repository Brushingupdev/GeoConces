"""Excel export for concessions using openpyxl."""

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.models import Concession


HEADERS = [
    ("code", "Código", 18),
    ("name", "Nombre", 36),
    ("holder_name", "Titular", 32),
    ("holder_ruc", "RUC", 14),
    ("status", "Estado", 12),
    ("concession_type", "Tipo", 16),
    ("region", "Región", 18),
    ("province", "Provincia", 18),
    ("district", "Distrito", 18),
    ("area_hectares", "Área (ha)", 12),
    ("registration_date", "Fecha registro", 16),
    ("title_date", "Fecha título", 16),
    ("expiration_date", "Fecha vencimiento", 18),
    ("source", "Fuente", 14),
]


def build_concessions_xlsx(concessions: Iterable[Concession]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Concesiones"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F6F69")
    header_align = Alignment(horizontal="left", vertical="center")

    for col_idx, (_, label, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, c in enumerate(concessions, start=2):
        for col_idx, (attr, _, _) in enumerate(HEADERS, start=1):
            value = getattr(c, attr, None)
            if hasattr(value, "isoformat"):
                value = value.strftime("%Y-%m-%d")
            elif hasattr(value, "value"):
                value = value.value
            ws.cell(row=row_idx, column=col_idx, value=value)

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
