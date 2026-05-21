from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import datetime
from typing import Any

from openpyxl import load_workbook


HEADER_ALIASES = {
    "codigo": "code",
    "codigounico": "code",
    "code": "code",
    "nombre": "name",
    "name": "name",
    "titular": "holder_name",
    "holder": "holder_name",
    "holdername": "holder_name",
    "ruc": "holder_ruc",
    "holderruc": "holder_ruc",
    "estado": "status",
    "status": "status",
    "tipo": "concession_type",
    "tipoconcesion": "concession_type",
    "concessiontype": "concession_type",
    "region": "region",
    "departamento": "region",
    "provincia": "province",
    "province": "province",
    "distrito": "district",
    "district": "district",
    "hectareas": "area_hectares",
    "areahectares": "area_hectares",
    "fecha_registro": "registration_date",
    "fecharegistro": "registration_date",
    "registrationdate": "registration_date",
    "fecha_titulo": "title_date",
    "fechatitulo": "title_date",
    "titledate": "title_date",
    "fecha_vencimiento": "expiration_date",
    "fechavencimiento": "expiration_date",
    "expirationdate": "expiration_date",
    "geometry": "geometry_wkt",
    "geom": "geometry_wkt",
    "geometrywkt": "geometry_wkt",
    "geojson": "geometry_geojson",
    "source": "source",
}


class IngestionParseError(ValueError):
    pass


def parse_tabular_upload(content: bytes, filename: str) -> list[dict[str, Any]]:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        return parse_csv_bytes(content)
    if lower_name.endswith(".xlsx"):
        return parse_xlsx_bytes(content)
    if lower_name.endswith(".geojson") or lower_name.endswith(".json"):
        return parse_geojson_bytes(content)
    if lower_name.endswith(".zip"):
        return parse_shapefile_zip_bytes(content)
    raise IngestionParseError("Unsupported file type. Use .csv, .xlsx, .geojson, or shapefile .zip")


def parse_csv_bytes(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return _normalize_rows(list(reader))


def parse_xlsx_bytes(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []

    raw_headers = [str(value).strip() if value is not None else "" for value in values[0]]
    rows: list[dict[str, Any]] = []
    for raw_row in values[1:]:
        if raw_row is None:
            continue
        row = {
            raw_headers[index]: raw_row[index]
            for index in range(len(raw_headers))
            if raw_headers[index]
        }
        rows.append(row)
    return _normalize_rows(rows)


def parse_geojson_bytes(content: bytes) -> list[dict[str, Any]]:
    payload = json.loads(content.decode("utf-8"))
    payload_type = payload.get("type")

    if payload_type == "FeatureCollection":
        features = payload.get("features", [])
    elif payload_type == "Feature":
        features = [payload]
    else:
        raise IngestionParseError("GeoJSON must be a Feature or FeatureCollection")

    rows: list[dict[str, Any]] = []
    for feature in features:
        properties = feature.get("properties") or {}
        row = dict(properties)
        row["geometry_geojson"] = feature.get("geometry")
        rows.append(row)

    return _normalize_rows(rows)


def parse_shapefile_zip_bytes(content: bytes) -> list[dict[str, Any]]:
    try:
        import shapefile
    except ModuleNotFoundError as exc:
        raise IngestionParseError("pyshp is required to parse shapefile uploads") from exc

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        lower_names = {name.lower(): name for name in archive.namelist()}
        shp_name = _find_archive_member(lower_names, ".shp")
        shx_name = _find_archive_member(lower_names, ".shx")
        dbf_name = _find_archive_member(lower_names, ".dbf")

        if not shp_name or not shx_name or not dbf_name:
            raise IngestionParseError("Shapefile zip must include .shp, .shx, and .dbf files")

        reader = shapefile.Reader(
            shp=io.BytesIO(archive.read(shp_name)),
            shx=io.BytesIO(archive.read(shx_name)),
            dbf=io.BytesIO(archive.read(dbf_name)),
        )

        rows: list[dict[str, Any]] = []
        field_names = [field[0] for field in reader.fields[1:]]
        for shape_record in reader.iterShapeRecords():
            attributes = dict(zip(field_names, shape_record.record))
            attributes["geometry_geojson"] = shape_record.shape.__geo_interface__
            rows.append(attributes)

    return _normalize_rows(rows)


def _normalize_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized_rows: list[dict[str, Any]] = []
    for row in rows:
        if not any(value not in (None, "") for value in row.values()):
            continue
        normalized = {
            _normalize_header(key): _normalize_cell(value)
            for key, value in row.items()
            if key is not None and _normalize_header(key) is not None
        }
        if not normalized.get("code"):
            raise IngestionParseError("Each row must include a code column")
        if not normalized.get("name"):
            normalized["name"] = normalized["code"]
        normalized_rows.append(normalized)
    return normalized_rows


def _find_archive_member(lower_names: dict[str, str], extension: str) -> str | None:
    for lower_name, original_name in lower_names.items():
        if lower_name.endswith(extension):
            return original_name
    return None


def _normalize_header(header: str | None) -> str | None:
    if header is None:
        return None
    compact = "".join(char for char in str(header).strip().lower() if char.isalnum() or char == "_")
    return HEADER_ALIASES.get(compact, compact or None)


def _normalize_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value
